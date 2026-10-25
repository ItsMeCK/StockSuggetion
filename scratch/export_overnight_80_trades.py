import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def export_overnight_80():
    df = pd.read_csv('data/grounded_201_trades_audit.csv')
    
    # Filter for overnight trades (where entry date != exit date)
    overnight = df[df['entry_time_ist'].str[:10] != df['exit_time_ist'].str[:10]].copy()
    overnight = overnight.reset_index(drop=True)
    
    # Re-index trade_no from 1 to 80
    overnight['trade_no'] = range(1, len(overnight) + 1)
    overnight['cumulative_net_pnl_inr'] = overnight['net_pnl_inr'].cumsum().round(2)
    
    csv_path = 'data/overnight_80_trades_audit.csv'
    excel_path = 'data/overnight_80_trades_audit.xlsx'
    
    overnight.to_csv(csv_path, index=False)
    print(f"✅ Exported CSV to: {csv_path}")

    # Write to Excel with formatting
    with pd.ExcelWriter(excel_path, engine='openpyxl') as writer:
        overnight.to_excel(writer, sheet_name='Overnight_80_Trades', index=False)
        workbook = writer.book
        worksheet = writer.sheets['Overnight_80_Trades']

        # Formatting Header
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        
        for col_idx in range(1, len(overnight.columns) + 1):
            cell = worksheet.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Color-code Net PnL and PnL %
        green_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
        green_font = Font(name="Calibri", size=11, color="276A3C", bold=True)
        red_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
        red_font = Font(name="Calibri", size=11, color="C00000", bold=True)

        pnl_col_idx = overnight.columns.get_loc('net_pnl_inr') + 1
        pct_col_idx = overnight.columns.get_loc('pnl_pct') + 1

        for row_idx in range(2, len(overnight) + 2):
            pnl_val = worksheet.cell(row=row_idx, column=pnl_col_idx).value
            if pnl_val is not None:
                if pnl_val > 0:
                    worksheet.cell(row=row_idx, column=pnl_col_idx).fill = green_fill
                    worksheet.cell(row=row_idx, column=pnl_col_idx).font = green_font
                    worksheet.cell(row=row_idx, column=pct_col_idx).fill = green_fill
                    worksheet.cell(row=row_idx, column=pct_col_idx).font = green_font
                elif pnl_val < 0:
                    worksheet.cell(row=row_idx, column=pnl_col_idx).fill = red_fill
                    worksheet.cell(row=row_idx, column=pnl_col_idx).font = red_font
                    worksheet.cell(row=row_idx, column=pct_col_idx).fill = red_fill
                    worksheet.cell(row=row_idx, column=pct_col_idx).font = red_font

        # Auto-adjust column widths
        for col in worksheet.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = col[0].column_letter
            worksheet.column_dimensions[col_letter].width = max(max_len + 3, 12)

    print(f"✅ Exported formatted Excel to: {excel_path}")

    # Print summary statistics
    total_net = overnight['net_pnl_inr'].sum()
    ce_net = overnight[overnight['option_type'] == 'CE']['net_pnl_inr'].sum()
    pe_net = overnight[overnight['option_type'] == 'PE']['net_pnl_inr'].sum()
    wins = len(overnight[overnight['net_pnl_inr'] > 0])
    losses = len(overnight[overnight['net_pnl_inr'] < 0])

    print("\n" + "=" * 60)
    print(f"Total Overnight Trades: {len(overnight)}")
    print(f"Win Rate: {wins/len(overnight)*100:.1f}% ({wins} W / {losses} L)")
    print(f"Total Net P&L: ₹{total_net:,.2f}")
    print(f"  - CE Net: ₹{ce_net:,.2f} ({len(overnight[overnight['option_type'] == 'CE'])} trades)")
    print(f"  - PE Net: ₹{pe_net:,.2f} ({len(overnight[overnight['option_type'] == 'PE'])} trades)")
    print("=" * 60)

if __name__ == '__main__':
    export_overnight_80()
